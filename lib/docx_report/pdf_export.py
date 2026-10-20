from __future__ import annotations

"""Export DOCX through Microsoft Word COM with placeholder protection."""

from pathlib import Path

from docx import Document


def _has_placeholders(docx: Path) -> bool:
    doc = Document(str(docx))
    return any("【图片占位" in p.text for p in doc.paragraphs) or any(
        "【图片占位" in p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
    )


def export_pdf(docx: Path, output: Path | None = None, *, allow_placeholders: bool = False,
               overwrite: bool = False) -> Path:
    docx = Path(docx).resolve()
    output = Path(output).resolve() if output else docx.with_suffix(".pdf")
    if not docx.exists():
        raise FileNotFoundError(docx)
    if output.exists() and not overwrite:
        raise FileExistsError(f"refusing to overwrite existing PDF: {output}")
    if not allow_placeholders and _has_placeholders(docx):
        raise RuntimeError("report still contains image placeholders; refusing final PDF export")

    import win32com.client  # imported lazily; Windows/Word only

    output.parent.mkdir(parents=True, exist_ok=True)
    word = win32com.client.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(docx), False, True, False)
        document.ExportAsFixedFormat(str(output), 17)
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()
    if not output.exists() or output.stat().st_size == 0:
        raise RuntimeError("Word returned without producing a non-empty PDF")
    return output
