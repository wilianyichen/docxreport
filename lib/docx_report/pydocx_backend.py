from __future__ import annotations

from dataclasses import dataclass

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


@dataclass(slots=True)
class ParagraphStyleSpec:
    bold: bool = False
    font_size: int = 12
    font_name: str = '宋体'
    space_after: int = 4
    first_line_indent: int | None = None
    alignment: WD_ALIGN_PARAGRAPH | None = None
    italic: bool = False


def replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    full = ''.join(run.text for run in paragraph.runs)
    new_full = full.replace(old, new)
    for index, run in enumerate(paragraph.runs):
        run.text = new_full if index == 0 else ''
    return True


def replace_in_cell(cell, old: str, new: str) -> bool:
    for paragraph in cell.paragraphs:
        if replace_in_paragraph(paragraph, old, new):
            return True
    return False


def add_paragraph(cell, text: str, spec: ParagraphStyleSpec | None = None):
    spec = spec or ParagraphStyleSpec()
    empty_paras = [p for p in cell.paragraphs if not p.text.strip()]
    paragraph = empty_paras[0] if empty_paras else cell.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = spec.bold
    run.italic = spec.italic
    run.font.size = Pt(spec.font_size)
    run.font.name = spec.font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), spec.font_name)
    pf = paragraph.paragraph_format
    pf.space_after = Pt(spec.space_after)
    if spec.first_line_indent:
        pf.first_line_indent = Pt(spec.first_line_indent)
    if spec.alignment is not None:
        paragraph.alignment = spec.alignment
    return paragraph


def add_heading(cell, text: str, *, font_size: int = 14):
    return add_paragraph(cell, text, ParagraphStyleSpec(bold=True, font_size=font_size, space_after=8))


def add_body(cell, text: str):
    return add_paragraph(cell, text, ParagraphStyleSpec(font_size=12, first_line_indent=24, space_after=4))


def add_placeholder(cell, text: str):
    return add_paragraph(cell, f'【截图占位 → {text}】', ParagraphStyleSpec(font_size=10, space_after=6))


def find_paragraph_after(doc: DocumentObject, keyword: str, skip: int = 0) -> int | None:
    found = 0
    for index, paragraph in enumerate(doc.paragraphs):
        if keyword in paragraph.text:
            if found >= skip:
                return index
            found += 1
    return None


def insert_image_after_paragraph(doc: DocumentObject, para_idx: int, image_path, *, width_inches: float = 5.5, caption: str = '') -> None:
    paragraph = doc.paragraphs[para_idx]
    if caption:
        cap_para = doc.add_paragraph()
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap_para.add_run(caption)
        run.font.size = Pt(9)
        run.italic = True
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        paragraph._element.addnext(cap_para._element)
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    if caption:
        cap_para._element.addnext(img_para._element)
    else:
        paragraph._element.addnext(img_para._element)
    spacer = doc.add_paragraph()
    img_para._element.addnext(spacer._element)
