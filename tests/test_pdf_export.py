"""PDF export guards. No Word is launched: every case fails before COM."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document

from docx_report.errors import (
    OverwriteRefusedError,
    PlaceholderRemainingError,
    PlatformNotSupportedError,
    ResourceMissingError,
)
from docx_report.pdf_export import PLACEHOLDER_MARKER, export_pdf, has_placeholders


def _docx_with(tmp_path, text: str):
    tmp_path = Path(tmp_path)
    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def test_placeholder_marker_detection(tmp_path):
    clean = _docx_with(tmp_path / "a", "正常正文")
    dirty = _docx_with(tmp_path / "b", f"{PLACEHOLDER_MARKER} → 运行截图】")

    assert has_placeholders(clean) is False
    assert has_placeholders(dirty) is True


def test_pdf_refuses_unreplaced_placeholders(tmp_path):
    path = _docx_with(tmp_path, f"{PLACEHOLDER_MARKER} → 运行截图】")
    with pytest.raises(PlaceholderRemainingError) as excinfo:
        export_pdf(path)
    assert excinfo.value.code == "E_PLACEHOLDER_REMAINING"


def test_pdf_refuses_to_overwrite_existing_file(tmp_path):
    path = _docx_with(tmp_path, "正常正文")
    existing = path.with_suffix(".pdf")
    existing.write_bytes(b"%PDF-1.4 stub")
    with pytest.raises(OverwriteRefusedError) as excinfo:
        export_pdf(path)
    assert excinfo.value.code == "E_OUTPUT_EXISTS"


def test_pdf_reports_missing_docx(tmp_path):
    with pytest.raises(ResourceMissingError):
        export_pdf(tmp_path / "nope.docx")


@pytest.mark.skipif(sys.platform != "win32", reason="Word COM is Windows-only")
def test_pdf_without_word_raises_platform_error_not_import_error(tmp_path):
    path = _docx_with(tmp_path, "正常正文")
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        with pytest.raises(PlatformNotSupportedError) as excinfo:
            export_pdf(path)
        assert excinfo.value.code == "E_PLATFORM_UNSUPPORTED"
        assert "pywin32" in str(excinfo.value)
        assert excinfo.value.hint
