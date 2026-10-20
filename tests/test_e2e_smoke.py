"""End-to-end smoke: build -> validate -> verify -> enrich, fully offline."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document

from docx_report.api import build_report, enrich_report, validate_report
from docx_report.verify import verify_document

CONFIG_TEMPLATE = """
name: e2e
template: template.docx
output: out/report.docx
cover_title: 实验七
experiment_name: 示例实验
blocks:
  - kind: section
    text: 一、实验目的
  - kind: body
    text: 掌握模板化报告的基本流程。
  - kind: section
    text: 二、实验原理
  - kind: figure
    image: fig1.png
    width_inches: 3.0
    text: 图 1 单图结果
  - kind: figure_group
    items:
      - image: fig1.png
        text: 图 2(a) 左图
      - image: fig2.png
        text: 图 2(b) 右图
  - kind: screenshot
    label: 运行截图
    hint: cell 3
    text: 图 3 单元运行截图
  - kind: caption
    text: 图 4 独立图注
  - kind: metrics
    source: results/metrics.json
    text: 结果：
    fields:
      - key: test.accuracy
        label: Accuracy
        fmt: ".3f"
"""

ENRICH_YAML = """
name: e2e-enrich
output: out/report-final.docx
actions:
  - action: add_page_break
  - action: append_heading
    text: 附录 A 运行实录
    level: 1
  - action: append_paragraph
    text: 以下为实际运行的终端输出。
  - action: append_code_from_file
    file: code.py
    head_lines: 3
  - action: append_image
    image: fig2.png
    width_inches: 3.0
  - action: insert_image_after_anchor
    anchor: 附录 A 运行实录
    image: fig1.png
    caption: 图 A-1 锚点插图
    width_inches: 2.5
"""


def _workspace(tmp_path: Path, template_docx: Path, png) -> Path:
    shutil.copy(template_docx, tmp_path / "template.docx")
    png("fig1.png")
    png("fig2.png", (120, 60))
    (tmp_path / "results").mkdir()
    (tmp_path / "results" / "metrics.json").write_text(
        json.dumps({"test": {"accuracy": 0.9234}}), encoding="utf-8"
    )
    (tmp_path / "code.py").write_text("import os\nprint(os.name)\n", encoding="utf-8")
    (tmp_path / "config.yml").write_text(CONFIG_TEMPLATE, encoding="utf-8")
    (tmp_path / "enrich.yml").write_text(ENRICH_YAML, encoding="utf-8")
    return tmp_path


def test_build_validate_verify(tmp_path, template_docx, png):
    work = _workspace(tmp_path, template_docx, png)
    output = build_report(work / "config.yml")

    assert output.exists() and output.stat().st_size > 0
    assert output == (work / "out" / "report.docx").resolve()

    validation = validate_report(work / "config.yml", output)
    assert validation.ok, validation.summary_lines()

    lines = verify_document(output, work / "config.yml")
    joined = "\n".join(lines)
    assert not any(line.startswith("ERROR") for line in lines)
    # figure + figure_group + screenshot => three "one image row + one caption row" tables
    assert "figure_group_tables: 3" in joined
    assert "placeholders: 0" in joined  # screenshot uses image placeholders, not text
    assert "missing_images: 0" in joined
    assert "page break after cover: YES" in joined

    doc = Document(str(output))
    text = "\n".join(
        p.text for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs
    )
    assert "一、实验目的" in text
    assert "Accuracy = 0.923" in text


def test_enrich_appends_and_anchors(tmp_path, template_docx, png):
    work = _workspace(tmp_path, template_docx, png)
    built = build_report(work / "config.yml")

    enriched = enrich_report(work / "enrich.yml", built)
    assert enriched.exists()

    validation = validate_report(work / "enrich.yml", enriched)
    assert validation.ok, validation.summary_lines()

    doc = Document(str(enriched))
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "附录 A 运行实录" in body_text
    assert "以下为实际运行的终端输出。" in body_text
    assert "import os" in body_text
    # anchor insert put an image somewhere after the heading
    assert any(p.text.strip() == "图 A-1 锚点插图" for p in doc.paragraphs)


def test_relative_paths_follow_the_config_directory(tmp_path, template_docx, png):
    work = _workspace(tmp_path, template_docx, png)
    nested = work / "nested" / "deep"
    nested.mkdir(parents=True)
    shutil.copy(work / "config.yml", nested / "config.yml")
    for name in ("template.docx", "fig1.png", "fig2.png"):
        shutil.copy(work / name, nested / name)
    shutil.copytree(work / "results", nested / "results")

    output = build_report(nested / "config.yml")
    assert output.exists()
    assert output.parent == (nested / "out").resolve()
