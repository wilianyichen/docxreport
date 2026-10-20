"""Export DOCX to PDF through Microsoft Word COM, with placeholder protection.

PDF export is intentionally the *last* step: image placeholders must be replaced
by hand first (Word/WPS "Change Picture" keeps the frame), so the exporter
refuses to run while any ``【图片占位`` text remains.

Platform support: Windows + Microsoft Word + pywin32 only. On other platforms a
``PlatformNotSupportedError`` (code ``E_PLATFORM_UNSUPPORTED``) is raised instead
of leaking a ``ModuleNotFoundError`` traceback.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

from .errors import (
    OverwriteRefusedError,
    PdfExportError,
    PlaceholderRemainingError,
    PlatformNotSupportedError,
    ResourceMissingError,
)

PLACEHOLDER_MARKER = "【图片占位"


def has_placeholders(docx: Path) -> bool:
    """True if the document still contains unreplaced image placeholders."""
    doc = Document(str(docx))
    if any(PLACEHOLDER_MARKER in p.text for p in doc.paragraphs):
        return True
    return any(
        PLACEHOLDER_MARKER in p.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )


# Backwards-compatible alias (older callers imported the private name).
_has_placeholders = has_placeholders


def _require_word():
    """Return the win32com.client module or raise a typed platform error."""
    if sys.platform != "win32":
        raise PlatformNotSupportedError(
            f"PDF export requires Microsoft Word on Windows (current platform: {sys.platform})",
            hint="export manually from Word/WPS, or run this step on a Windows machine",
        )
    try:
        import win32com.client  # noqa: F401 - imported lazily, Windows/Word only
    except ImportError as exc:  # pragma: no cover - depends on host
        raise PlatformNotSupportedError(
            "pywin32 is not installed; PDF export needs it to drive Microsoft Word",
            hint='pip install "docxreport[pdf]"',
        ) from exc
    return win32com.client


def export_pdf(
    docx: Path, output: Path | None = None, *, allow_placeholders: bool = False, overwrite: bool = False
) -> Path:
    """Convert ``docx`` to PDF via Word. Returns the written PDF path."""
    docx = Path(docx).resolve()
    output = Path(output).resolve() if output else docx.with_suffix(".pdf")
    if not docx.exists():
        raise ResourceMissingError(f"docx not found: {docx}")
    if output.exists() and not overwrite:
        raise OverwriteRefusedError(
            f"refusing to overwrite existing PDF: {output}", hint="pass --overwrite to replace it"
        )
    if not allow_placeholders and has_placeholders(docx):
        raise PlaceholderRemainingError(
            "report still contains image placeholders; refusing final PDF export",
            hint="replace every placeholder image first, or pass --allow-placeholders for a draft",
        )

    win32com_client = _require_word()
    output.parent.mkdir(parents=True, exist_ok=True)
    word = win32com_client.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(docx), False, True, False)
        document.ExportAsFixedFormat(str(output), 17)
    except Exception as exc:  # noqa: BLE001 - COM raises bare Exception subclasses
        raise PdfExportError(f"Word failed to export PDF: {exc}", context={"docx": str(docx)}) from exc
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()
    if not output.exists() or output.stat().st_size == 0:
        raise PdfExportError(
            "Word returned without producing a non-empty PDF", context={"output": str(output)}
        )
    return output
