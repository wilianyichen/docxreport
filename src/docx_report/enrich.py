from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .config_loader import load_enrichment_config
from .errors import AnchorNotFoundError, EnrichmentActionError, ResourceMissingError
from .models import EnrichmentConfig
from .pydocx_backend import find_paragraph_after, insert_image_after_paragraph

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def enrich_existing_report(config_or_path, input_docx: Path, output_path: Path | None = None) -> Path:
    config = (
        config_or_path
        if isinstance(config_or_path, EnrichmentConfig)
        else load_enrichment_config(Path(config_or_path))
    )
    doc = Document(str(input_docx))
    for action in config.actions:
        _apply_action(doc, config, action)
    final_output = output_path or config.output or input_docx
    doc.save(str(final_output))
    return final_output


def _apply_action(doc: Document, config: EnrichmentConfig, action) -> None:
    if action.action == "insert_image_after_anchor":
        if not action.anchor or not action.image:
            raise EnrichmentActionError(
                "insert_image_after_anchor requires anchor and image",
                hint="set both `anchor:` and `image:` on the action",
            )
        required = bool(action.extras.get("required", True))
        image_path = _resolve_path(config, action.image)
        if not image_path.exists():
            if required:
                raise ResourceMissingError(f"image not found: {image_path}")
            return
        idx = find_paragraph_after(doc, action.anchor, skip=int(action.extras.get("skip", 0) or 0))
        if idx is None:
            if required:
                raise AnchorNotFoundError(
                    f"anchor not found: {action.anchor}",
                    hint="check the heading text; run `doctor` to see which anchors exist",
                )
            return
        insert_image_after_paragraph(
            doc,
            idx,
            image_path,
            width_inches=action.width_inches or 5.5,
            caption=action.caption,
        )
        return

    if action.action == "add_page_break":
        doc.add_page_break()
        return

    if action.action == "append_heading":
        text = str(action.extras.get("text") or "")
        if not text:
            raise EnrichmentActionError("append_heading requires text")
        _append_heading(doc, text, level=int(action.extras.get("level", 1) or 1))
        return

    if action.action == "append_paragraph":
        text = str(action.extras.get("text") or "")
        if not text:
            raise EnrichmentActionError("append_paragraph requires text")
        _append_paragraph(
            doc,
            text,
            bold=bool(action.extras.get("bold", False)),
            font_size=float(action.extras.get("font_size", 12) or 12),
            font_name=str(action.extras.get("font_name") or "宋体"),
            alignment=_parse_alignment(action.extras.get("alignment")),
        )
        return

    if action.action == "append_code_lines":
        lines = _coerce_lines(action.extras.get("lines"))
        if not lines:
            raise EnrichmentActionError("append_code_lines requires lines")
        _append_code_lines(doc, lines, font_size=float(action.extras.get("font_size", 7.5) or 7.5))
        return

    if action.action == "append_code_from_file":
        raw_path = action.extras.get("file")
        if not raw_path:
            raise EnrichmentActionError("append_code_from_file requires file")
        file_path = _resolve_path(config, Path(str(raw_path)))
        lines = _read_head_lines(file_path, int(action.extras.get("head_lines", 0) or 0))
        tab_columns = action.extras.get("tab_columns")
        if tab_columns is not None:
            lines = _truncate_tabs(lines, int(tab_columns))
        lines.extend(_coerce_lines(action.extras.get("append_lines")))
        _append_code_lines(doc, lines, font_size=float(action.extras.get("font_size", 7.5) or 7.5))
        return

    if action.action == "append_image":
        if not action.image:
            raise EnrichmentActionError("append_image requires image")
        required = bool(action.extras.get("required", True))
        image_path = _resolve_path(config, action.image)
        if not image_path.exists():
            if required:
                raise ResourceMissingError(f"image not found: {image_path}")
            return
        _append_image(doc, image_path, width_inches=action.width_inches or 5.5, caption=action.caption)
        return

    raise EnrichmentActionError(f"unsupported enrichment action: {action.action}")


def _append_heading(doc: Document, text: str, *, level: int) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        _set_run_style(run, font_name="宋体")


def _append_paragraph(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    font_size: float = 12,
    font_name: str = "宋体",
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> None:
    paragraph = doc.add_paragraph()
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    _set_run_style(run, font_name=font_name, font_size=font_size, bold=bold)


def _append_code_lines(doc: Document, lines: list[str], *, font_size: float = 7.5) -> None:
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line)
        _set_run_style(run, font_name="Courier New", font_size=font_size)
    doc.add_paragraph()


def _append_image(doc: Document, image_path: Path, *, width_inches: float = 5.5, caption: str = "") -> None:
    if caption:
        caption_paragraph = doc.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(caption)
        run.italic = True
        _set_run_style(run, font_name="宋体", font_size=9)
    doc.add_picture(str(image_path), width=Inches(width_inches))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def _set_run_style(run, *, font_name: str, font_size: float = 12, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _parse_alignment(value) -> WD_ALIGN_PARAGRAPH | None:
    if value is None:
        return None
    alignment = ALIGNMENTS.get(str(value).lower())
    if alignment is None:
        raise ValueError(f"unsupported alignment: {value}")
    return alignment


def _resolve_path(config: EnrichmentConfig, path: Path) -> Path:
    if path.is_absolute() or config.source is None:
        return path
    return config.source.parent / path


def _read_head_lines(path: Path, head_lines: int) -> list[str]:
    if not path.exists():
        raise ResourceMissingError(f"artifact file not found: {path}")
    if head_lines <= 0:
        head_lines = 10
    lines: list[str] = []
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for _, line in zip(range(head_lines), handle, strict=False):
            lines.append(line.rstrip("\n"))
    return lines


def _truncate_tabs(lines: list[str], max_columns: int) -> list[str]:
    truncated: list[str] = []
    for line in lines:
        fields = line.split("\t")
        if len(fields) > max_columns:
            truncated.append("\t".join(fields[:max_columns]) + "\t...")
        else:
            truncated.append(line)
    return truncated


def _coerce_lines(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value]
    return [str(value)]
