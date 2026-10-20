"""Helpers for writing tests against docx_report (also used by `examples/`).

Public so downstream projects can generate the same fixture template instead of
copying a course document around:

    from docx_report.testing import build_minimal_template
    template = build_minimal_template(tmp_path / "template.docx")
"""

from __future__ import annotations

from pathlib import Path

__all__ = ["build_minimal_template"]


def build_minimal_template(path: str | Path, *, donor_text: str = "一、实验目的") -> Path:
    """Write the smallest template that satisfies the fixed-table assumptions.

    Layout:

    * body paragraph ``实验一`` carrying the page break (cover title slot)
    * table 0: cover information table (1 row, 2 columns)
    * table 1: report table (6 rows, 2 columns)
      - row 3 / col 1: experiment name slot
      - row 5 / col 0: content slot, seeded with a style donor paragraph
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Pt

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.styles["Normal"].font.name = "SimSun"
    doc.styles["Normal"].font.size = Pt(12)

    title = doc.add_paragraph("实验一")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(18)
    # The real templates keep the page break INSIDE the cover title paragraph;
    # update_cover_title() must preserve it (verify: "page break after cover").
    title.runs[0].add_break(WD_BREAK.PAGE)

    cover = doc.add_table(rows=1, cols=2)
    cover.cell(0, 0).text = "课程名称"
    cover.cell(0, 1).text = "示例课程"

    report = doc.add_table(rows=6, cols=2)
    report.cell(0, 0).text = "实验报告"
    report.cell(3, 0).text = "实验名称"
    report.cell(3, 1).text = "【待填写实验名称】"

    donor = report.cell(5, 0).paragraphs[0]
    donor_run = donor.add_run(donor_text)
    donor_run.font.name = "SimSun"
    donor_run.font.size = Pt(12)
    donor.paragraph_format.line_spacing = 1.5
    # python-docx only writes w:ascii/w:hAnsi; set the CJK slot so the donor
    # carries a complete w:rFonts triple (ascii/hAnsi/eastAsia).
    rfonts = donor_run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")

    doc.save(str(target))
    return target
