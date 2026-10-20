from __future__ import annotations

import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document

from .config_loader import _load_mapping, load_enrichment_config, load_report_config
from .experiment_readiness import check_experiment_readiness
from .models import DoctorCheck, DoctorResult, EnrichmentConfig, ReportConfig
from .ooxml_backend import (
    find_cover_title_paragraph,
    read_document_xml,
    register_namespaces,
    resolve_template_slots,
    w_tag,
)
from .pydocx_backend import find_paragraph_after
from .style_resolver import resolve_style_profile
from .validators import validate_report_file, validate_template_layout


def entry_command() -> str:
    """Portable invocation string used in "next step" hints.

    Order matters: prefer the installed console script (most readable), fall back
    to this repo's stable shim, then to ``python -m``. Resolved at runtime so the
    hint is always copy-pasteable on the current machine.
    """
    console = shutil.which("docx-report")
    if console:
        return "docx-report"
    repo_root = Path(__file__).resolve().parents[2]
    entry = repo_root / "scripts" / "skill_entry.py"
    if entry.exists():
        return f'"{sys.executable}" "{entry}"'
    return "python -m docx_report"


SKILL_ENTRY = entry_command()


def doctor_config(config_or_path, input_docx: Path | None = None) -> DoctorResult:
    input_path = Path(input_docx) if input_docx is not None else None
    if isinstance(config_or_path, ReportConfig):
        return _doctor_report_config(config_or_path, input_path)
    if isinstance(config_or_path, EnrichmentConfig):
        return _doctor_enrichment_config(config_or_path, input_path)

    config_path = Path(config_or_path)
    raw = _load_mapping(config_path)
    if isinstance(raw.get("actions"), list) and not raw.get("blocks"):
        return _doctor_enrichment_config(load_enrichment_config(config_path), input_path)
    return _doctor_report_config(load_report_config(config_path), input_path)


def _doctor_report_config(config: ReportConfig, input_docx: Path | None) -> DoctorResult:
    checks: list[DoctorCheck] = []
    notes: list[str] = []
    ok = True

    checks.append(DoctorCheck("ok", "config", f"loaded report config: {config.name}"))
    checks.append(DoctorCheck("ok", "blocks", f"{len(config.blocks)} blocks loaded"))

    placeholder_count = sum(1 for block in config.blocks if block.kind == "placeholder")
    checks.append(DoctorCheck("ok", "placeholders", f"{placeholder_count} placeholder blocks"))

    caption_count = sum(1 for block in config.blocks if block.kind == "caption")
    if placeholder_count or caption_count:
        if caption_count == placeholder_count:
            checks.append(
                DoctorCheck(
                    "ok",
                    "captions",
                    f"{caption_count} caption blocks (1:1 with {placeholder_count} placeholders)",
                )
            )
        else:
            checks.append(
                DoctorCheck(
                    "warn",
                    "captions",
                    f"{caption_count} caption blocks vs {placeholder_count} placeholders (not 1:1)",
                )
            )

    figure_blocks = [b for b in config.blocks if b.kind == "figure"]
    if figure_blocks:
        base = config.source.parent if config.source is not None else Path.cwd()
        missing = []
        for b in figure_blocks:
            raw = str(b.meta.get("image") or "")
            if not raw:
                missing.append("<empty image path>")
                continue
            cand = Path(raw)
            if not cand.is_absolute():
                cand = base / cand
            if not cand.exists():
                missing.append(raw)
        if missing:
            checks.append(
                DoctorCheck("error", "figures", f"{len(missing)} figure image(s) not found: {missing[:3]}")
            )
            ok = False
        else:
            checks.append(
                DoctorCheck("ok", "figures", f"{len(figure_blocks)} figure blocks, all image files exist")
            )

    profile_path = (getattr(config, "extras", None) or {}).get("_course_profile_resolved")
    if profile_path:
        checks.append(DoctorCheck("ok", "course-profile", f"course report profile applied: {profile_path}"))
    else:
        notes.append(
            "no course report profile reached (kind: course-report-profile); using built-in defaults"
        )

    output_parent = config.output.parent
    if output_parent.exists():
        checks.append(DoctorCheck("ok", "output", f"output directory exists: {output_parent}"))
    else:
        checks.append(
            DoctorCheck(
                "warn", "output", f"output directory missing now, will be created on write: {output_parent}"
            )
        )

    template_ok = False
    if config.template is None:
        ok = False
        checks.append(DoctorCheck("error", "template", "report config has no template; build is unavailable"))
    else:
        checks.append(DoctorCheck("ok", "template", f"template path set: {config.template}"))
        template_result = validate_template_layout(config.template, config.assumptions)
        template_ok = template_result.ok
        ok = ok and template_result.ok
        for message in template_result.messages:
            status = "ok" if template_result.ok else "error"
            checks.append(DoctorCheck(status, "template-layout", message))

    template_body = (
        _load_report_body(config.template)
        if config.template is not None and config.template.exists()
        else None
    )
    template_cell = (
        _load_report_content_cell_from_body(template_body, config) if template_body is not None else None
    )
    template_title_paragraph = (
        find_cover_title_paragraph(template_body) if template_body is not None else None
    )
    template_experiment_name_cell = (
        _load_report_name_cell_from_body(template_body, config) if template_body is not None else None
    )

    if input_docx is not None:
        report_result = validate_report_file(input_docx)
        ok = ok and report_result.ok
        for message in report_result.messages:
            status = "ok" if report_result.ok else "error"
            checks.append(DoctorCheck(status, "input-docx", message))
        input_body = _load_report_body(input_docx) if report_result.ok else None
        input_cell = (
            _load_report_content_cell_from_body(input_body, config) if input_body is not None else None
        )
        input_title_paragraph = find_cover_title_paragraph(input_body) if input_body is not None else None
        input_experiment_name_cell = (
            _load_report_name_cell_from_body(input_body, config) if input_body is not None else None
        )
        if report_result.ok and config.expected_images is not None:
            notes.append(f"revise mode expects {config.expected_images} preserved images")
    else:
        input_body = None
        input_cell = None
        input_title_paragraph = None
        input_experiment_name_cell = None
        notes.append("no input_docx provided; doctor can assess build readiness but not revise readiness")

    resolved = resolve_style_profile(
        config,
        template_cell=template_cell,
        input_cell=input_cell,
        template_title_paragraph=template_title_paragraph,
        input_title_paragraph=input_title_paragraph,
        template_experiment_name_cell=template_experiment_name_cell,
        input_experiment_name_cell=input_experiment_name_cell,
    )
    for kind in ("cover_title", "experiment_name", "section", "subheading", "body", "placeholder"):
        status = "error" if kind in resolved.unresolved_required else "ok"
        suffix = " (required donor unresolved)" if kind in resolved.unresolved_required else ""
        checks.append(
            DoctorCheck(
                status,
                f"style-source:{kind}",
                f"{kind} style source: {resolved.sources.get(kind, 'fallback')}{suffix}",
            )
        )
    if resolved.skipped_comment_count:
        checks.append(
            DoctorCheck(
                "ok",
                "style-source:comments",
                f"skipped {resolved.skipped_comment_count} comment/placeholder donor paragraphs",
            )
        )

    requirements_raw = (config.extras or {}).get("experiment_requirements")
    if requirements_raw:
        req_path = Path(str(requirements_raw))
        if not req_path.is_absolute():
            req_path = (config.source.parent if config.source else Path.cwd()) / req_path
        try:
            readiness = check_experiment_readiness(req_path)
            for item in readiness.items:
                detail = item.message if len(item.message) <= 160 else item.message[:157] + "..."
                status = "ok" if item.status in {"complete", "optional_not_done"} else "error"
                checks.append(DoctorCheck(status, "experiment:" + item.item_id, item.status + " - " + detail))
            if not readiness.ok:
                ok = False
                notes.append("experiment gate blocked: fix mandatory items before reporting them as done")
        except Exception as exc:
            ok = False
            checks.append(DoctorCheck("error", "experiment-readiness", str(exc)))

    config_ref = str(config.source or config.name)
    if input_docx is not None and input_docx.exists():
        suggested_command = f"{SKILL_ENTRY} revise {config_ref} {input_docx}"
    elif template_ok:
        suggested_command = f"{SKILL_ENTRY} build {config_ref}"
    else:
        suggested_command = f"{SKILL_ENTRY} inspect-template {config.template}" if config.template else None

    return DoctorResult(
        ok=ok,
        config_type="report",
        config_name=config.name,
        checks=checks,
        suggested_command=suggested_command,
        notes=notes,
    )


def _doctor_enrichment_config(config: EnrichmentConfig, input_docx: Path | None) -> DoctorResult:
    checks: list[DoctorCheck] = []
    notes: list[str] = []
    ok = True

    checks.append(DoctorCheck("ok", "config", f"loaded enrich config: {config.name}"))
    checks.append(DoctorCheck("ok", "actions", f"{len(config.actions)} actions loaded"))

    output_parent = config.output.parent if config.output is not None else None
    if output_parent is None:
        checks.append(
            DoctorCheck(
                "warn",
                "output",
                "config has no output; enrich will overwrite the input report unless --output is provided",
            )
        )
    elif output_parent.exists():
        checks.append(DoctorCheck("ok", "output", f"output directory exists: {output_parent}"))
    else:
        checks.append(
            DoctorCheck(
                "warn", "output", f"output directory missing now, will be created on write: {output_parent}"
            )
        )

    anchor_count = 0
    for index, action in enumerate(config.actions, start=1):
        label = f"action[{index}]"
        if action.action == "insert_image_after_anchor":
            anchor_count += 1
            if not action.anchor:
                ok = False
                checks.append(DoctorCheck("error", label, "insert_image_after_anchor is missing anchor"))
            else:
                checks.append(DoctorCheck("ok", label, f"anchor declared: {action.anchor}"))
            if action.image is None:
                ok = False
                checks.append(DoctorCheck("error", label, "insert_image_after_anchor is missing image"))
            else:
                ok = ok and _check_resource_exists(
                    checks,
                    config,
                    action.image,
                    label,
                    required=bool(action.extras.get("required", True)),
                    purpose="anchor image",
                )
            continue

        if action.action == "append_image":
            if action.image is None:
                ok = False
                checks.append(DoctorCheck("error", label, "append_image is missing image"))
            else:
                ok = ok and _check_resource_exists(
                    checks,
                    config,
                    action.image,
                    label,
                    required=bool(action.extras.get("required", True)),
                    purpose="append image",
                )
            continue

        if action.action == "append_code_from_file":
            raw_file = action.extras.get("file")
            if not raw_file:
                ok = False
                checks.append(DoctorCheck("error", label, "append_code_from_file is missing file"))
            else:
                ok = ok and _check_resource_exists(
                    checks, config, Path(str(raw_file)), label, required=True, purpose="artifact file"
                )
            continue

    if anchor_count == 0:
        notes.append("this enrich config has no anchor-based insertions; it only appends content/resources")

    if input_docx is not None:
        report_result = validate_report_file(input_docx)
        ok = ok and report_result.ok
        for message in report_result.messages:
            status = "ok" if report_result.ok else "error"
            checks.append(DoctorCheck(status, "input-docx", message))
        if report_result.ok:
            doc = Document(str(input_docx))
            for index, action in enumerate(config.actions, start=1):
                if action.action != "insert_image_after_anchor" or not action.anchor:
                    continue
                label = f"anchor[{index}]"
                idx = find_paragraph_after(doc, action.anchor, skip=int(action.extras.get("skip", 0) or 0))
                if idx is None:
                    ok = False
                    checks.append(
                        DoctorCheck("error", label, f"anchor not found in input report: {action.anchor}")
                    )
                else:
                    checks.append(DoctorCheck("ok", label, f"anchor found in input report: {action.anchor}"))
    else:
        notes.append("no input_docx provided; doctor can assess resource readiness but not anchor presence")

    config_ref = str(config.source or config.name)
    if input_docx is not None:
        suggested_command = f"{SKILL_ENTRY} enrich {config_ref} {input_docx}"
    else:
        suggested_command = f"{SKILL_ENTRY} enrich {config_ref} <input.docx>"

    return DoctorResult(
        ok=ok,
        config_type="enrich",
        config_name=config.name,
        checks=checks,
        suggested_command=suggested_command,
        notes=notes,
    )


def _load_report_body(docx_path: Path):
    document_xml, _, _ = read_document_xml(docx_path)
    register_namespaces(document_xml)
    root = ET.fromstring(document_xml)
    return root.find(w_tag("body"))


def _load_report_content_cell_from_body(body, config: ReportConfig):
    if body is None:
        return None
    try:
        _, _, content_cell = resolve_template_slots(body, config)
    except Exception:
        return None
    return content_cell


def _load_report_name_cell_from_body(body, config: ReportConfig):
    if body is None:
        return None
    try:
        _, report_name_cell, _ = resolve_template_slots(body, config)
    except Exception:
        return None
    return report_name_cell


def _check_resource_exists(
    checks: list[DoctorCheck],
    config: EnrichmentConfig,
    path: Path,
    label: str,
    *,
    required: bool,
    purpose: str,
) -> bool:
    resolved = _resolve_enrichment_path(config, path)
    if resolved.exists():
        checks.append(DoctorCheck("ok", label, f"{purpose} exists: {resolved}"))
        return True
    if required:
        checks.append(DoctorCheck("error", label, f"{purpose} missing: {resolved}"))
        return False
    checks.append(DoctorCheck("warn", label, f"optional {purpose} missing: {resolved}"))
    return True


def _resolve_enrichment_path(config: EnrichmentConfig, path: Path) -> Path:
    if path.is_absolute() or config.source is None:
        return path
    return config.source.parent / path
