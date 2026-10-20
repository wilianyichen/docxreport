"""交付核验：把「人工读一遍 docx」变成一条命令的结构化摘要。

用途：替代把整份报告读回上下文来核对，显著降低 token 消耗。
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

TEMPLATE_LEFTOVER_KEYWORDS = (
    "结合实验指导书",
    "不能直接复制粘贴",
    "要求：按作业练习逐项展开",
    "注意1：截图",
    "注意2：结果数据",
    "归纳总结",
    "【说明",
    "（要求：",
    "（以上四部分",
)

SECTION_PREFIXES = ("一、", "二、", "三、", "四、", "五、", "六、")


def _fmt(paragraph) -> str:
    parts = []
    pPr = paragraph.find(qn("w:pPr"))
    if pPr is not None:
        sp = pPr.find(qn("w:spacing"))
        if sp is not None and sp.get(qn("w:line")):
            parts.append(f"line={sp.get(qn('w:line'))}")
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            parts.append(f"align={jc.get(qn('w:val'))}")
    r = paragraph.find(qn("w:r"))
    if r is not None:
        rPr = r.find(qn("w:rPr"))
        if rPr is not None:
            sz = rPr.find(qn("w:sz"))
            if sz is not None:
                parts.append(f"sz={sz.get(qn('w:val'))}")
            parts.append("bold" if rPr.find(qn("w:b")) is not None else "regular")
            color = rPr.find(qn("w:color"))
            if color is not None:
                parts.append(f"color={color.get(qn('w:val'))}")
    return " ".join(parts)


def _cover_title_page_break(doc) -> tuple[bool, str]:
    """模板封面标题段落（形如「实验一」）必须保留 <w:br w:type="page"/>。

    update_cover_title() 历史上会重建该段落并丢掉分页符，导致封面与正文挤在同一页。
    """
    import re

    pattern = re.compile(r"^实验(?:[0-9]+|[一二三四五六七八九十]+|\*+)$")
    for paragraph in doc.element.body.findall(qn("w:p")):
        text = "".join(n.text or "" for n in paragraph.iter(qn("w:t"))).strip()
        if not pattern.match(text):
            continue
        breaks = paragraph.findall(".//" + qn("w:br"))
        return any(br.get(qn("w:type")) == "page" for br in breaks), text
    return False, ""


def _content_cell(doc, assumptions=None):
    """Locate the content cell.

    Preferred: the config's own ``assumptions`` (table/row/col indices), so any
    template layout works. Fallback: the historical keyword heuristic
    (a cell containing both 实验目的 and 实验原理).
    """
    if assumptions is not None:
        try:
            table = doc.tables[assumptions.report_table_index]
            return table.rows[assumptions.content_row].cells[assumptions.content_col]._tc
        except (IndexError, KeyError):
            pass
    for table in doc.element.body.findall(qn("w:tbl")):
        for row in table.findall(qn("w:tr")):
            for cell in row.findall(qn("w:tc")):
                text = "".join(n.text or "" for n in cell.iter(qn("w:t")))
                if "实验目的" in text and "实验原理" in text:
                    return cell
    return None


def verify_document(docx_path: Path, config_path: Path | None = None) -> list[str]:
    docx_path = Path(docx_path)
    lines: list[str] = []
    if not docx_path.exists():
        return [f"ERROR: file not found: {docx_path}"]
    lines.append(f"file: {docx_path.name} ({docx_path.stat().st_size:,} bytes)")

    cfg = None
    if config_path is not None:
        try:
            from .config_loader import load_report_config

            cfg = load_report_config(Path(config_path))
        except Exception:  # noqa: BLE001 - verify must still work without a valid config
            cfg = None

    try:
        doc = Document(str(docx_path))
    except Exception as exc:  # noqa: BLE001
        return lines + [f"ERROR: cannot open with python-docx: {exc}"]

    paragraphs = [p for p in doc.element.body.iter() if p.tag == qn("w:p")]
    full_text = "\n".join("".join(n.text or "" for n in p.iter(qn("w:t"))) for p in paragraphs)
    cn = sum(1 for ch in full_text if "\u4e00" <= ch <= "\u9fff")
    lines.append(f"paragraphs: {len(paragraphs)} | chinese_chars: {cn}")

    cell = _content_cell(doc, cfg.assumptions if cfg is not None else None)
    if cell is None:
        lines.append("WARN: content cell not located (pass the config, or use 实验目的/实验原理 headings)")
    else:
        cell_paras = [p for p in cell.iter() if p.tag == qn("w:p")]
        imgs = [p for p in cell_paras if p.findall(".//" + qn("w:drawing"))]
        body_imgs = [p for p in paragraphs if p.findall(".//" + qn("w:drawing"))]
        lines.append(f"content_cell_paragraphs: {len(cell_paras)}")
        lines.append(
            f"images_in_content_cell: {len(imgs)} | images_elsewhere: {len(body_imgs) - len(imgs)}"
            "   (elsewhere 通常为封面校徽/页眉图，属正常)"
        )

    lines.append("\n[section headings]")
    for p in paragraphs:
        text = "".join(n.text or "" for n in p.iter(qn("w:t"))).strip()
        if text.startswith(SECTION_PREFIXES):
            lines.append(f"  {text}  ->  {_fmt(p)}")

    lines.append("\n[figure blocks]")
    caption_count = 0
    placeholder_count = 0
    for p in paragraphs:
        text = "".join(n.text or "" for n in p.iter(qn("w:t"))).strip()
        if text.startswith("【图片占位"):
            placeholder_count += 1
        elif text.startswith("图") or text.startswith("运行截图"):
            caption_count += 1
    lines.append(f"  placeholders: {placeholder_count} | captions: {caption_count}")
    nested = cell.findall(qn("w:tbl")) if cell is not None else []
    if nested:
        images_in_tables = sum(len(t.findall(".//" + qn("w:drawing"))) for t in nested)
        two_row = sum(1 for t in nested if len(t.findall(qn("w:tr"))) == 2)
        lines.append(
            f"  figure_group_tables: {len(nested)} (2-row: {two_row}) | images_in_tables: {images_in_tables}"
        )

    has_break, title_text = _cover_title_page_break(doc)
    lines.append("\n[cover page break]")
    lines.append(f"  cover_title paragraph: {title_text!r}")
    if has_break:
        lines.append("  page break after cover: YES")
    else:
        lines.append("  page break after cover: NO  <== WARN 封面与正文会挤在同一页（模板分页符被丢弃）")

    lines.append("\n[template leftovers]")
    leftovers = [kw for kw in TEMPLATE_LEFTOVER_KEYWORDS if kw in full_text]
    lines.append("  none" if not leftovers else f"  FOUND: {leftovers}")

    try:
        with zipfile.ZipFile(docx_path) as zf:
            media = [n for n in zf.namelist() if n.startswith("word/media/") and not n.endswith("/")]
        lines.append(f"\n[package] media_files: {len(media)}")
    except Exception as exc:  # noqa: BLE001
        lines.append(f"\n[package] WARN: {exc}")

    if config_path is not None:
        if cfg is None:
            lines.append("\n[config] WARN: config could not be loaded")
            return lines
        lines.append("\n[config]")
        lines.append(f"  course_profile: {cfg.extras.get('_course_profile_resolved')}")
        figures = [b for b in cfg.blocks if b.kind == "figure"]
        base = cfg.source.parent if cfg.source is not None else Path.cwd()
        missing = []
        for b in figures:
            raw = str(b.meta.get("image") or "")
            cand = Path(raw)
            if not cand.is_absolute():
                cand = base / cand
            if not raw or not cand.exists():
                missing.append(raw)
        lines.append(f"  figure_blocks: {len(figures)} | missing_images: {len(missing)}")
    return lines
